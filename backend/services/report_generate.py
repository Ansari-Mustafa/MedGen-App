"""
Report generation service — renders DOCX template with filled JSON.
Logic ported from extras/scripts/generate_report.py.
Returns the generated DOCX as bytes.
"""
import asyncio
import copy
import os
import re
import tempfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

LIST_FIELD_RE = re.compile(r"\{\{-?\s*(\w+)\[\]\s*-?\}\}")


def _para_full_text(para_elem) -> str:
    return "".join(t.text or "" for t in para_elem.iter(qn("w:t")))


def _set_para_text(para_elem, text: str, remove_list_style: bool = False):
    for tag in [qn("w:r"), qn("w:hyperlink"), qn("w:ins"), qn("w:del")]:
        for child in para_elem.findall(tag):
            para_elem.remove(child)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    para_elem.append(r)
    if remove_list_style:
        pPr = para_elem.find(qn("w:pPr"))
        if pPr is not None:
            for tag in [qn("w:pStyle"), qn("w:numPr")]:
                el = pPr.find(tag)
                if el is not None:
                    pPr.remove(el)


def _transform_list_placeholders(root_elem) -> int:
    to_transform = []
    for para_el in root_elem.findall(".//" + qn("w:p")):
        m = LIST_FIELD_RE.search(_para_full_text(para_el))
        if m:
            to_transform.append((para_el, m.group(1)))

    for para_el, field_name in to_transform:
        parent = para_el.getparent()
        if parent is None:
            continue
        idx = list(parent).index(para_el)

        for_elem = copy.deepcopy(para_el)
        _set_para_text(for_elem, "{%p for item in " + field_name + " %}", remove_list_style=True)
        _set_para_text(para_el, "{{ item }}", remove_list_style=False)
        endfor_elem = copy.deepcopy(para_el)
        _set_para_text(endfor_elem, "{%p endfor %}", remove_list_style=True)

        parent.insert(idx, for_elem)
        parent.insert(idx + 2, endfor_elem)

    return len(to_transform)


def _normalise_fill_payload(data: dict) -> tuple[dict, list[dict]]:
    """
    Accept either the legacy wrapped shape `{filled: {...}, unfilled: [...]}`
    or the flat placeholder→value map the pipeline now passes.
    """
    if isinstance(data, dict) and "filled" in data and isinstance(data["filled"], dict):
        return data["filled"], data.get("unfilled") or []
    return (data or {}), []


def _build_context(data: dict, keep_unfilled: bool = True) -> dict:
    filled, unfilled = _normalise_fill_payload(data)
    unfilled_keys = {item["placeholder"] for item in unfilled if isinstance(item, dict) and "placeholder" in item}
    context: dict = {}
    for k, v in filled.items():
        is_list = isinstance(v, list)
        if v is None:
            if keep_unfilled:
                context[k] = ["{{ " + k + "[] }}"] if is_list else "{{" + k + "}}"
            else:
                context[k] = [] if is_list else ""
        else:
            context[k] = v
    for key in unfilled_keys:
        if key not in context:
            context[key] = "{{" + key + "}}" if keep_unfilled else ""
    return context


def _render(template_bytes: bytes, data: dict, keep_unfilled: bool = True) -> bytes:
    """Synchronous render — runs in a thread."""
    # Write template to temp file
    fd_tpl, tpl_path = tempfile.mkstemp(suffix=".docx")
    os.close(fd_tpl)

    fd_out, out_path = tempfile.mkstemp(suffix=".docx")
    os.close(fd_out)

    try:
        with open(tpl_path, "wb") as f:
            f.write(template_bytes)

        # Pre-process list placeholders
        doc = Document(tpl_path)
        count = _transform_list_placeholders(doc.element.body)

        processed_path = tpl_path
        if count > 0:
            fd_proc, processed_path = tempfile.mkstemp(suffix=".docx")
            os.close(fd_proc)
            doc.save(processed_path)

        tpl = DocxTemplate(processed_path)
        tpl.render(_build_context(data, keep_unfilled))
        tpl.save(out_path)

        if processed_path != tpl_path:
            os.unlink(processed_path)

        return open(out_path, "rb").read()

    finally:
        for p in (tpl_path, out_path):
            try:
                os.unlink(p)
            except FileNotFoundError:
                pass


async def generate_report(template_bytes: bytes, filled_data: dict, keep_unfilled: bool = True) -> bytes:
    """
    Render a DOCX template with filled_data.
    Returns the rendered DOCX as bytes.
    """
    return await asyncio.to_thread(_render, template_bytes, filled_data, keep_unfilled)
